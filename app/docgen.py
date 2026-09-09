"""
Final stage: render the generated section content into a polished .docx.

Kept intentionally simple (one shared layout for every doc_type) rather than
building a bespoke template per document type - see the README tradeoff
discussion ("Simplicity vs Extensibility").
"""
from __future__ import annotations

import os
import re
from datetime import datetime
from typing import Dict, List

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Inches

OUTPUT_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "output")
os.makedirs(OUTPUT_DIR, exist_ok=True)

ACCENT = RGBColor(0x1F, 0x4E, 0x79)


def _slugify(text: str) -> str:
    slug = re.sub(r"[^a-zA-Z0-9]+", "_", text.strip().lower()).strip("_")
    return slug[:60] or "document"


def _add_bulleted_or_paragraph(doc: Document, content: str) -> None:
    lines = [l.strip() for l in content.splitlines() if l.strip()]
    bullet_lines = [l for l in lines if l.startswith(("-", "*", "•"))]
    if bullet_lines and len(bullet_lines) == len(lines):
        for line in lines:
            text = line.lstrip("-*• ").strip()
            doc.add_paragraph(text, style="List Bullet")
    else:
        # Preserve any inline bullet sub-lines within an otherwise prose section
        buffer: List[str] = []
        for line in lines:
            if line.startswith(("-", "*", "•")):
                if buffer:
                    doc.add_paragraph(" ".join(buffer))
                    buffer = []
                doc.add_paragraph(line.lstrip("-*• ").strip(), style="List Bullet")
            else:
                buffer.append(line)
        if buffer:
            doc.add_paragraph(" ".join(buffer))


def render_docx(
    title: str,
    doc_type: str,
    sections: List[str],
    section_content: Dict[str, str],
    assumptions: List[str],
    generation_mode: str,
) -> str:
    doc = Document()

    # ---- Title page / header block ----
    heading = doc.add_heading(title, level=0)
    for run in heading.runs:
        run.font.color.rgb = ACCENT

    meta = doc.add_paragraph()
    meta.add_run(f"Document type: {doc_type.replace('_', ' ').title()}").italic = True
    meta.add_run(f"  |  Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}").italic = True
    meta.alignment = WD_ALIGN_PARAGRAPH.LEFT

    doc.add_paragraph(
        "Generated autonomously by the AI Agent Document Builder."
        + (" (LLM-assisted content)" if generation_mode == "llm" else " (fallback template mode - LLM unavailable)")
    ).italic = True

    doc.add_paragraph()  # spacer

    # ---- Assumptions callout (only if any were made) ----
    if assumptions:
        doc.add_heading("Assumptions Made by the Agent", level=1)
        doc.add_paragraph(
            "The request did not fully specify every detail. To complete the task "
            "autonomously, the agent made the following reasonable assumptions:"
        )
        for a in assumptions:
            doc.add_paragraph(a, style="List Bullet")
        doc.add_paragraph()

    # ---- Body sections ----
    for section in sections:
        doc.add_heading(section, level=1)
        content = section_content.get(section, "[No content generated]")
        _add_bulleted_or_paragraph(doc, content)

    filename = f"{_slugify(doc_type)}_{_slugify(title)}_{int(datetime.now().timestamp())}.docx"
    path = os.path.join(OUTPUT_DIR, filename)
    doc.save(path)
    return path
